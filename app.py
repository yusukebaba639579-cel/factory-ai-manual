from __future__ import annotations

import json
import sqlite3
import uuid
from datetime import datetime, timezone
from pathlib import Path
from urllib.request import Request, urlopen
from urllib.request import urlretrieve

import cv2
import mediapipe as mp
import numpy as np
from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader
from mediapipe.tasks import python
from mediapipe.tasks.python import vision
from flask import Flask, jsonify, redirect, render_template, request, send_file, url_for
from werkzeug.utils import secure_filename

BASE_DIR = Path(__file__).resolve().parent
UPLOAD_DIR = BASE_DIR / "uploads"
DATABASE = BASE_DIR / "factory_ai_manual.db"
MODEL_DIR = BASE_DIR / "models"
POSE_MODEL = MODEL_DIR / "pose_landmarker_lite.task"
HAND_MODEL = MODEL_DIR / "hand_landmarker.task"
POSE_URL = "https://storage.googleapis.com/mediapipe-models/pose_landmarker/pose_landmarker_lite/float16/latest/pose_landmarker_lite.task"
HAND_URL = "https://storage.googleapis.com/mediapipe-models/hand_landmarker/hand_landmarker/float16/latest/hand_landmarker.task"
UPLOAD_DIR.mkdir(exist_ok=True)
MODEL_DIR.mkdir(exist_ok=True)

app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = 2 * 1024 * 1024 * 1024


def connect():
    db = sqlite3.connect(DATABASE)
    db.row_factory = sqlite3.Row
    db.execute("PRAGMA foreign_keys = ON")
    return db


def init_database():
    with connect() as db:
        db.executescript("""
        CREATE TABLE IF NOT EXISTS videos (
          id INTEGER PRIMARY KEY AUTOINCREMENT,
          title TEXT NOT NULL, file_name TEXT NOT NULL, file_path TEXT NOT NULL,
          duration REAL NOT NULL DEFAULT 0, rotation INTEGER NOT NULL DEFAULT 0,
          procedure_name TEXT NOT NULL DEFAULT '', viewer_analysis_version INTEGER NOT NULL DEFAULT 0,
          created_at TEXT NOT NULL
        );
        CREATE TABLE IF NOT EXISTS processes (
          id INTEGER PRIMARY KEY AUTOINCREMENT,
          video_id INTEGER NOT NULL REFERENCES videos(id) ON DELETE CASCADE,
          position INTEGER NOT NULL, title TEXT NOT NULL,
          start_time REAL NOT NULL, end_time REAL NOT NULL,
          description_ja TEXT NOT NULL DEFAULT '', feature_json TEXT NOT NULL DEFAULT '[]',
          UNIQUE(video_id, position)
        );
        CREATE TABLE IF NOT EXISTS translations (
          id INTEGER PRIMARY KEY AUTOINCREMENT,
          process_id INTEGER NOT NULL REFERENCES processes(id) ON DELETE CASCADE,
          language_code TEXT NOT NULL, title TEXT NOT NULL DEFAULT '', description TEXT NOT NULL,
          UNIQUE(process_id, language_code)
        );
        CREATE TABLE IF NOT EXISTS training_examples (
          id INTEGER PRIMARY KEY AUTOINCREMENT,
          label TEXT NOT NULL, feature_json TEXT NOT NULL, created_at TEXT NOT NULL
        );
        """)
        columns = {row[1] for row in db.execute("PRAGMA table_info(processes)")}
        if "feature_json" not in columns:
            db.execute("ALTER TABLE processes ADD COLUMN feature_json TEXT NOT NULL DEFAULT '[]'")
        translation_columns = {row[1] for row in db.execute("PRAGMA table_info(translations)")}
        if "title" not in translation_columns:
            db.execute("ALTER TABLE translations ADD COLUMN title TEXT NOT NULL DEFAULT ''")
        video_columns = {row[1] for row in db.execute("PRAGMA table_info(videos)")}
        if "rotation" not in video_columns:
            db.execute("ALTER TABLE videos ADD COLUMN rotation INTEGER NOT NULL DEFAULT 0")
        if "procedure_name" not in video_columns:
            db.execute("ALTER TABLE videos ADD COLUMN procedure_name TEXT NOT NULL DEFAULT ''")
        if "viewer_analysis_version" not in video_columns:
            db.execute("ALTER TABLE videos ADD COLUMN viewer_analysis_version INTEGER NOT NULL DEFAULT 0")
        if "active_json" not in columns:
            db.execute("ALTER TABLE processes ADD COLUMN active_json TEXT NOT NULL DEFAULT ''")
        if "focus_x" not in columns:
            db.execute("ALTER TABLE processes ADD COLUMN focus_x REAL NOT NULL DEFAULT 0.5")
        if "focus_y" not in columns:
            db.execute("ALTER TABLE processes ADD COLUMN focus_y REAL NOT NULL DEFAULT 0.5")


init_database()


def ensure_model(path: Path, url: str) -> str:
    if not path.exists():
        try:
            urlretrieve(url, path)
        except OSError as error:
            path.unlink(missing_ok=True)
            raise ValueError("MediaPipeモデルを取得できません。初回のみネット接続が必要です。") from error
    return str(path)


def detect_process_ranges(path: Path, rotation: int = 0) -> tuple[float, list[dict]]:
    """Find a variable number of boundaries from image, arm and hand motion."""
    capture = cv2.VideoCapture(str(path))
    if not capture.isOpened():
        raise ValueError("動画を読み込めませんでした。MP4をお試しください。")
    fps = capture.get(cv2.CAP_PROP_FPS) or 30.0
    frames = int(capture.get(cv2.CAP_PROP_FRAME_COUNT) or 0)
    duration = frames / fps if frames else 0
    if duration < 3:
        capture.release()
        raise ValueError("3秒以上の動画を使用してください。")
    pose_options = vision.PoseLandmarkerOptions(base_options=python.BaseOptions(model_asset_path=ensure_model(POSE_MODEL, POSE_URL)), running_mode=vision.RunningMode.VIDEO, num_poses=1, min_pose_detection_confidence=.25, min_tracking_confidence=.25)
    hand_options = vision.HandLandmarkerOptions(base_options=python.BaseOptions(model_asset_path=ensure_model(HAND_MODEL, HAND_URL)), running_mode=vision.RunningMode.VIDEO, num_hands=2, min_hand_detection_confidence=.25, min_tracking_confidence=.25)
    interval = 0.5
    samples, previous_gray, previous_body = [], None, None
    with vision.PoseLandmarker.create_from_options(pose_options) as pose, vision.HandLandmarker.create_from_options(hand_options) as hands:
        for time in np.arange(0, duration, interval):
            capture.set(cv2.CAP_PROP_POS_MSEC, float(time * 1000))
            ok, frame = capture.read()
            if not ok:
                continue
            if rotation == 180:
                frame = cv2.rotate(frame, cv2.ROTATE_180)
            gray = cv2.resize(cv2.cvtColor(frame, cv2.COLOR_BGR2GRAY), (48, 27))
            image_motion = 0.0 if previous_gray is None else float(np.mean(cv2.absdiff(gray, previous_gray)) / 255)
            image = mp.Image(image_format=mp.ImageFormat.SRGB, data=cv2.cvtColor(frame, cv2.COLOR_BGR2RGB))
            timestamp = int(time * 1000)
            pose_result, hand_result = pose.detect_for_video(image, timestamp), hands.detect_for_video(image, timestamp)
            body = []
            if pose_result.pose_landmarks:
                for index in (11, 12, 13, 14, 15, 16):
                    point = pose_result.pose_landmarks[0][index]
                    body.extend((point.x, point.y))
            body.extend(value for hand in hand_result.hand_landmarks for point in hand for value in (point.x, point.y))
            body_array = np.asarray(body[:84], dtype=float)
            body_motion = 0.0
            if previous_body is not None and len(body_array) == len(previous_body) and len(body_array):
                body_motion = float(np.mean(np.abs(body_array - previous_body)))
            hand_count = len(hand_result.hand_landmarks)
            samples.append({"time": float(time), "image": image_motion, "body": body_motion, "hands": hand_count})
            previous_gray, previous_body = gray, body_array
    capture.release()
    if len(samples) < 5:
        return duration, [{"start": 0.0, "end": duration, "feature": [0, 0, 0]}]
    raw = np.asarray([s["image"] + s["body"] * 5 for s in samples])
    smooth = np.convolve(raw, np.ones(3) / 3, mode="same")
    change = np.abs(np.diff(smooth, prepend=smooth[0]))
    threshold = max(float(np.percentile(change, 75)), .008)
    minimum_gap, boundaries, last = 2.0, [0.0], 0.0
    for index in np.argsort(change)[::-1]:
        time = samples[int(index)]["time"]
        if change[index] < threshold or time < minimum_gap or duration - time < minimum_gap:
            continue
        if all(abs(time - boundary) >= minimum_gap for boundary in boundaries):
            boundaries.append(round(time, 2))
            last = time
        if len(boundaries) >= 12:
            break
    boundaries.append(round(duration, 2))
    boundaries.sort()
    segments = []
    for start, end in zip(boundaries, boundaries[1:]):
        selected = [s for s in samples if start <= s["time"] < end]
        if end - start < 1.5 or not selected:
            continue
        feature = [float(np.mean([s["image"] for s in selected])), float(np.mean([s["body"] for s in selected])), float(np.mean([s["hands"] for s in selected]))]
        segments.append({"start": start, "end": end, "feature": feature})
    return duration, segments or [{"start": 0.0, "end": duration, "feature": [0, 0, 0]}]


def analyze_viewer_activity(path: Path, processes, rotation: int = 0) -> list[dict]:
    """Locate useful motion and its visual center without modifying the source video."""
    capture = cv2.VideoCapture(str(path))
    if not capture.isOpened():
        return []
    results = []
    for process in processes:
        start, end, interval = float(process["start_time"]), float(process["end_time"]), 0.16
        samples, previous = [], None
        for timestamp in np.arange(start, end, interval):
            capture.set(cv2.CAP_PROP_POS_MSEC, float(timestamp * 1000))
            ok, frame = capture.read()
            if not ok:
                continue
            if rotation == 180:
                frame = cv2.rotate(frame, cv2.ROTATE_180)
            gray = cv2.resize(cv2.cvtColor(frame, cv2.COLOR_BGR2GRAY), (160, 90))
            if previous is None:
                samples.append((float(timestamp), 0.0, 0.5, 0.5))
            else:
                difference = cv2.absdiff(gray, previous).astype(np.float32) / 255
                score = float(np.mean(difference))
                weights = np.maximum(difference - 0.055, 0)
                total = float(weights.sum())
                if total:
                    yy, xx = np.indices(weights.shape)
                    focus_x = float((xx * weights).sum() / total / weights.shape[1])
                    focus_y = float((yy * weights).sum() / total / weights.shape[0])
                else:
                    focus_x, focus_y = 0.5, 0.5
                samples.append((float(timestamp), score, focus_x, focus_y))
            previous = gray
        if len(samples) < 3:
            results.append({"id": process["id"], "active": [[start, end]], "focus_x": 0.5, "focus_y": 0.5})
            continue
        scores = np.asarray([sample[1] for sample in samples])
        threshold = max(0.008, float(np.percentile(scores, 58)) * 0.85)
        active_indexes = [index for index, sample in enumerate(samples) if sample[1] >= threshold]
        if not active_indexes:
            active_indexes = [int(np.argmax(scores))]
        ranges = []
        group_start = group_end = active_indexes[0]
        for index in active_indexes[1:]:
            if index - group_end <= 2:
                group_end = index
            else:
                ranges.append([max(start, samples[group_start][0] - 0.2), min(end, samples[group_end][0] + interval + 0.2)])
                group_start = group_end = index
        ranges.append([max(start, samples[group_start][0] - 0.2), min(end, samples[group_end][0] + interval + 0.2)])
        merged = []
        for item in ranges:
            if merged and item[0] - merged[-1][1] <= 0.3:
                merged[-1][1] = max(merged[-1][1], item[1])
            else:
                merged.append(item)
        ranges = merged
        if sum(item[1] - item[0] for item in ranges) < min(0.8, end - start):
            peak = samples[int(np.argmax(scores))][0]
            ranges = [[max(start, peak - 0.45), min(end, peak + 0.45)]]
        motion_samples = [samples[index] for index in active_indexes]
        weights = np.asarray([max(item[1], 0.001) for item in motion_samples])
        focus_x = float(np.average([item[2] for item in motion_samples], weights=weights))
        focus_y = float(np.average([item[3] for item in motion_samples], weights=weights))
        results.append({"id": process["id"], "active": [[round(a, 2), round(b, 2)] for a, b in ranges], "focus_x": round(min(0.78, max(0.22, focus_x)), 3), "focus_y": round(min(0.78, max(0.22, focus_y)), 3)})
    capture.release()
    return results


def ensure_viewer_analysis(video, processes) -> None:
    if not video or not processes or (int(video["viewer_analysis_version"] or 0) >= 2 and all(row["active_json"] for row in processes)):
        return
    path = Path(video["file_path"])
    if not path.exists():
        return
    results = analyze_viewer_activity(path, processes, int(video["rotation"] or 0))
    with connect() as db:
        for result in results:
            db.execute("UPDATE processes SET active_json=?,focus_x=?,focus_y=? WHERE id=?", (json.dumps(result["active"]), result["focus_x"], result["focus_y"], result["id"]))
        db.execute("UPDATE videos SET viewer_analysis_version=2 WHERE id=?", (video["id"],))


def extract_procedure_text(uploaded) -> tuple[str, str]:
    if not uploaded or not uploaded.filename:
        return "", ""
    extension = Path(uploaded.filename).suffix.lower()
    if extension not in {".pdf", ".docx", ".xlsx", ".csv", ".txt"}:
        raise ValueError("手順書はPDF・Word・Excel・CSV・TXTに対応しています。")
    safe_name = secure_filename(uploaded.filename) or f"procedure{extension}"
    path = UPLOAD_DIR / f"{uuid.uuid4().hex}_{safe_name}"
    uploaded.save(path)
    try:
        if extension == ".pdf":
            text = "\n".join(page.extract_text() or "" for page in PdfReader(path).pages)
        elif extension == ".docx":
            text = "\n".join(paragraph.text for paragraph in Document(path).paragraphs)
        elif extension == ".xlsx":
            workbook = load_workbook(path, read_only=True, data_only=True)
            text = "\n".join(" | ".join(str(value) for value in row if value is not None) for sheet in workbook.worksheets for row in sheet.iter_rows(values_only=True))
        else:
            text = path.read_text(encoding="utf-8-sig", errors="ignore")
    except Exception as error:
        raise ValueError(f"手順書を読み取れませんでした: {error}") from error
    if not text.strip():
        raise ValueError("手順書から文字を読み取れませんでした。")
    return text[:16000], uploaded.filename


def fit_segments_to_steps(segments: list[dict], duration: float, count: int) -> list[dict]:
    count = max(1, min(count, 30))
    candidates = [item["end"] for item in segments[:-1]]
    boundaries = [0.0]
    for index in range(1, count):
        target = duration * index / count
        usable = [value for value in candidates if value > boundaries[-1] + .7 and duration - value > (count - index) * .7]
        boundaries.append(min(usable, key=lambda value: abs(value - target)) if usable else target)
    boundaries.append(duration)
    fitted = []
    for start, end in zip(boundaries, boundaries[1:]):
        midpoint = (start + end) / 2
        source = min(segments, key=lambda item: abs((item["start"] + item["end"]) / 2 - midpoint))
        fitted.append({"start": round(start, 2), "end": round(end, 2), "feature": source["feature"]})
    return fitted


def generate_from_procedure(manual_title: str, procedure_text: str) -> list[dict]:
    prompt = f"""次の社内手順書を動画マニュアル用の作業一覧へ変換してください。
手順書に書かれた順番・作業数・作業名を最優先してください。注意書きだけの行は独立作業にしないでください。
各説明は日本語1文、30文字以内。JSONのみを返してください。
形式: {{"processes":[{{"title":"作業名","description":"短い説明"}}]}}
マニュアル名: {manual_title}
手順書:\n{procedure_text}"""
    try:
        data = json.dumps({"model": "phi4", "prompt": prompt, "stream": False, "format": "json", "keep_alive": "10m", "options": {"temperature": .1, "num_predict": 900}}).encode("utf-8")
        req = Request("http://127.0.0.1:11434/api/generate", data=data, headers={"Content-Type": "application/json"})
        with urlopen(req, timeout=180) as response:
            result = json.loads(json.loads(response.read())["response"])["processes"]
    except Exception as error:
        raise ValueError("Phi-4で手順書を解析できませんでした。") from error
    clean = [{"title": str(item.get("title", "作業")).strip()[:40], "description": str(item.get("description", "")).strip()[:30]} for item in result if str(item.get("title", "")).strip()]
    if not clean:
        raise ValueError("手順書から作業を抽出できませんでした。")
    return clean


def predict_labels(segments: list[dict]) -> list[str]:
    allowed = ["ねじ締め", "ねじの締付確認", "不適合確認"]
    with connect() as db:
        examples = db.execute("SELECT label, feature_json FROM training_examples").fetchall()
    learned = [(row["label"], np.asarray(json.loads(row["feature_json"]))) for row in examples if row["label"] in allowed]
    labels = []
    activity = [float(item["feature"][0]) + float(item["feature"][1]) * 5 for item in segments]
    median_activity = float(np.median(activity)) if activity else 0
    for segment in segments:
        feature = np.asarray(segment["feature"])
        if len(learned) >= 6:
            nearest = sorted(((float(np.linalg.norm(feature - saved)), label) for label, saved in learned))[:3]
            labels.append(max(allowed, key=lambda label: sum(1 / (distance + .001) for distance, found in nearest if found == label)))
        else:
            movement = float(feature[0]) + float(feature[1]) * 5
            hands = float(feature[2])
            if hands >= .5 and movement >= median_activity:
                labels.append("ねじ締め")
            elif hands >= .5 or movement >= median_activity * .65:
                labels.append("ねじの締付確認")
            else:
                labels.append("不適合確認")
    return labels


def generate_processes_with_phi4(manual_title: str, segments: list[dict]) -> list[dict]:
    labels = predict_labels(segments)
    timing = [{"number": i + 1, "start": item["start"], "end": item["end"], "seconds": round(item["end"] - item["start"], 1), "label": labels[i]} for i, item in enumerate(segments)]
    prompt = f"""あなたは製造現場の標準作業書を作る専門家です。
マニュアル名: {manual_title}
OpenCVが映像変化から検出した作業候補: {json.dumps(timing, ensure_ascii=False)}
作業名は候補のlabelをそのまま使ってください。labelは「ねじ締め」「ねじの締付確認」「不適合確認」「判定保留」のいずれかです。
作業数を変えないでください。各作業の説明を日本語で1文、30文字以内で作ってください。推測した固有部品名や数値は書かないでください。
JSONのみを返してください。形式: {{"processes":[{{"number":1,"title":"...","description":"..."}}]}}"""
    try:
        data = json.dumps({"model": "phi4", "prompt": prompt, "stream": False, "format": "json", "keep_alive": "10m", "options": {"temperature": 0.1, "num_predict": 180}}).encode("utf-8")
        req = Request("http://127.0.0.1:11434/api/generate", data=data, headers={"Content-Type": "application/json"})
        with urlopen(req, timeout=180) as response:
            result = json.loads(json.loads(response.read())["response"])["processes"]
    except Exception as error:
        raise ValueError("Ollama + Phi-4で作業文を生成できませんでした。Ollamaを起動し、ollama run phi4 を実行してください。") from error
    if not isinstance(result, list) or len(result) != len(segments):
        raise ValueError("Phi-4の生成結果を作業データとして読み取れませんでした。もう一度お試しください。")
    for index, item in enumerate(result):
        item["title"] = labels[index]
        description = str(item.get("description", "")).strip().replace("\n", " ")
        if not description:
            raise ValueError("Phi-4が説明文を生成できませんでした。もう一度お試しください。")
        item["description"] = description[:30]
    return result


@app.get("/")
def home():
    with connect() as db:
        videos = db.execute("SELECT v.*, COUNT(p.id) process_count FROM videos v LEFT JOIN processes p ON p.video_id=v.id GROUP BY v.id ORDER BY v.id DESC").fetchall()
    return render_template("index.html", videos=videos)


@app.post("/api/videos")
def create_video():
    uploaded = request.files.get("video")
    if not uploaded or not uploaded.filename:
        return jsonify(error="動画を選択してください。"), 400
    extension = Path(uploaded.filename).suffix.lower()
    if extension not in {".mp4", ".mov", ".webm", ".mkv", ".avi"}:
        return jsonify(error="MP4 / MOV / WebM / MKV / AVIに対応しています。"), 400
    safe_name = secure_filename(uploaded.filename) or f"video{extension}"
    stored = UPLOAD_DIR / f"{uuid.uuid4().hex}_{safe_name}"
    uploaded.save(stored)
    title = request.form.get("title", "").strip() or Path(uploaded.filename).stem
    rotation = 180 if request.form.get("rotate") == "180" else 0
    try:
        procedure_text, procedure_name = extract_procedure_text(request.files.get("procedure"))
    except ValueError as error:
        stored.unlink(missing_ok=True)
        return jsonify(error=str(error)), 400
    try:
        duration, ranges = detect_process_ranges(stored, rotation)
    except ValueError:
        capture = cv2.VideoCapture(str(stored))
        fps, frames = capture.get(cv2.CAP_PROP_FPS) or 30, capture.get(cv2.CAP_PROP_FRAME_COUNT) or 0
        duration = frames / fps if frames else 0
        capture.release()
        if duration <= 0:
            stored.unlink(missing_ok=True)
            return jsonify(error="動画を読み込めませんでした。別のMP4動画をお試しください。"), 400
        ranges = [{"start": 0.0, "end": round(duration, 2), "feature": [0, 0, 0]}]
    if procedure_text:
        try:
            generated = generate_from_procedure(title, procedure_text)
            ranges = fit_segments_to_steps(ranges, duration, len(generated))
        except ValueError:
            generated = [{"title": "手順書確認", "description": "手順書を確認して作業内容を登録してください。"}]
            ranges = fit_segments_to_steps(ranges, duration, 1)
    else:
        try:
            generated = generate_processes_with_phi4(title, ranges)
        except ValueError:
            generated = [{"title": "判定保留", "description": "映像を確認して作業内容を登録してください。"} for _ in ranges]
    with connect() as db:
        cursor = db.execute("INSERT INTO videos(title,file_name,file_path,duration,rotation,procedure_name,created_at) VALUES(?,?,?,?,?,?,?)", (title, uploaded.filename, str(stored), duration, rotation, procedure_name, datetime.now(timezone.utc).isoformat()))
        video_id = cursor.lastrowid
        for position, (segment, text) in enumerate(zip(ranges, generated), 1):
            db.execute("INSERT INTO processes(video_id,position,title,start_time,end_time,description_ja,feature_json) VALUES(?,?,?,?,?,?,?)", (video_id, position, str(text.get("title", "判定保留")).strip(), segment["start"], segment["end"], str(text.get("description", "")).strip(), json.dumps(segment["feature"])))
    return jsonify(id=video_id, manual_url=f"/manuals/{video_id}")


def load_manual(video_id):
    with connect() as db:
        video = db.execute("SELECT * FROM videos WHERE id=?", (video_id,)).fetchone()
        processes = db.execute("SELECT * FROM processes WHERE video_id=? ORDER BY position", (video_id,)).fetchall()
    return video, processes


@app.get("/manuals/<int:video_id>/edit")
def edit_manual(video_id):
    video, processes = load_manual(video_id)
    if not video:
        return "マニュアルが見つかりません。", 404
    return render_template("editor.html", video=video, processes=processes)


@app.get("/manuals/<int:video_id>")
def view_manual(video_id):
    language_code = request.args.get("lang", "ja")
    if language_code not in {"ja", "en", "vi", "zh"}:
        language_code = "ja"
    video, base_processes = load_manual(video_id)
    if not video:
        return "マニュアルが見つかりません。", 404
    ensure_viewer_analysis(video, base_processes)
    with connect() as db:
        if language_code == "ja":
            processes = db.execute("SELECT * FROM processes WHERE video_id=? ORDER BY position", (video_id,)).fetchall()
        else:
            processes = db.execute("""
                SELECT p.id, p.video_id, p.position, p.start_time, p.end_time,
                       p.feature_json, p.active_json, p.focus_x, p.focus_y,
                       COALESCE(NULLIF(t.title, ''), p.title) AS title,
                       COALESCE(NULLIF(t.description, ''), p.description_ja) AS description_ja
                FROM processes AS p
                LEFT JOIN translations AS t
                  ON t.process_id = p.id AND t.language_code = ?
                WHERE p.video_id = ?
                ORDER BY p.position
            """, (language_code, video_id)).fetchall()
    return render_template("manual.html", video=video, processes=processes, language_code=language_code)


def ensure_argos_pair(source_code: str, target_code: str) -> None:
    from argostranslate import package, translate

    try:
        if translate.get_translation_from_codes(source_code, target_code) is not None:
            return
    except Exception:
        pass
    package.update_package_index()
    available = package.get_available_packages()
    direct = next((item for item in available if item.from_code == source_code and item.to_code == target_code), None)
    if direct:
        package.install_from_path(direct.download())
        translate.get_installed_languages.cache_clear()
        return
    # Argos can pivot automatically. Install Japanese→English and
    # English→target when a direct package is unavailable.
    pairs = [(source_code, "en"), ("en", target_code)]
    for from_code, to_code in pairs:
        try:
            if translate.get_translation_from_codes(from_code, to_code) is not None:
                continue
        except Exception:
            pass
        model = next((item for item in available if item.from_code == from_code and item.to_code == to_code), None)
        if model is None:
            raise ValueError(f"Argos翻訳モデル {from_code}→{to_code} が見つかりません。")
        package.install_from_path(model.download())
        translate.get_installed_languages.cache_clear()


def argos_translate_text(text: str, target_code: str) -> str:
    from argostranslate import translate

    ensure_argos_pair("ja", target_code)
    translate.get_installed_languages.cache_clear()
    try:
        translator = translate.get_translation_from_codes("ja", target_code)
    except (AttributeError, StopIteration):
        translator = None
    if translator is not None:
        return translator.translate(text)
    # Explicit pivot is a fallback for Argos versions that do not rebuild the
    # composite Japanese→English→target route immediately after installation.
    first = translate.get_translation_from_codes("ja", "en")
    second = translate.get_translation_from_codes("en", target_code)
    if first is None or second is None:
        raise ValueError("インストール済み翻訳モデルを読み込めません。")
    return second.translate(first.translate(text))


@app.post("/manuals/<int:video_id>/translate")
def translate_manual(video_id):
    language_code = request.form.get("language", "")
    languages = {"en": "英語", "vi": "ベトナム語", "zh": "中国語"}
    if language_code not in languages:
        return redirect(url_for("view_manual", video_id=video_id))
    video, processes = load_manual(video_id)
    if not video or not processes:
        return "翻訳する作業がありません。", 400
    try:
        ensure_argos_pair("ja", language_code)
        translated = [{"id": row["id"], "title": argos_translate_text(row["title"], language_code), "description": argos_translate_text(row["description_ja"], language_code)} for row in processes]
    except Exception as error:
        return f"Argos Translateで翻訳できませんでした: {error}", 503
    with connect() as db:
        for item in translated:
            db.execute("INSERT INTO translations(process_id,language_code,title,description) VALUES(?,?,?,?) ON CONFLICT(process_id,language_code) DO UPDATE SET title=excluded.title,description=excluded.description", (int(item["id"]), language_code, str(item["title"]).strip(), str(item["description"]).strip()))
    return redirect(url_for("view_manual", video_id=video_id, lang=language_code))


@app.get("/api/videos/<int:video_id>/stream")
def stream_video(video_id):
    with connect() as db:
        video = db.execute("SELECT file_path FROM videos WHERE id=?", (video_id,)).fetchone()
    if not video or not Path(video["file_path"]).exists():
        return jsonify(error="動画が見つかりません。"), 404
    return send_file(video["file_path"], conditional=True)


@app.post("/api/videos/<int:video_id>/processes")
def create_process(video_id):
    payload = request.get_json(silent=True) or {}
    try:
        title = str(payload["title"]).strip()
        start, end = round(float(payload["start_time"]), 2), round(float(payload["end_time"]), 2)
        if not title or start < 0 or end <= start:
            raise ValueError
    except (KeyError, TypeError, ValueError):
        return jsonify(error="作業名と正しい開始・終了位置を入力してください。"), 400
    with connect() as db:
        video = db.execute("SELECT duration FROM videos WHERE id=?", (video_id,)).fetchone()
        if not video:
            return jsonify(error="動画が見つかりません。"), 404
        if video["duration"] and end > video["duration"] + .5:
            return jsonify(error="終了位置が動画の長さを超えています。"), 400
        position = db.execute("SELECT COALESCE(MAX(position),0)+1 FROM processes WHERE video_id=?", (video_id,)).fetchone()[0]
        cursor = db.execute("INSERT INTO processes(video_id,position,title,start_time,end_time,description_ja) VALUES(?,?,?,?,?,?)", (video_id, position, title, start, end, f"{title}を行います。作業後に状態を確認してください。"))
    return jsonify(id=cursor.lastrowid, position=position)


@app.patch("/api/processes/<int:process_id>")
def update_process(process_id):
    payload = request.get_json(silent=True) or {}
    title, description = str(payload.get("title", "")).strip(), str(payload.get("description_ja", "")).strip()
    if not title:
        return jsonify(error="作業名を入力してください。"), 400
    with connect() as db:
        current = db.execute("SELECT title,feature_json FROM processes WHERE id=?", (process_id,)).fetchone()
        if not current:
            return jsonify(error="作業が見つかりません。"), 404
        cursor = db.execute("UPDATE processes SET title=?,description_ja=? WHERE id=?", (title, description, process_id))
        if title in {"ねじ締め", "ねじの締付確認", "不適合確認"} and current["feature_json"] != "[]":
            db.execute("INSERT INTO training_examples(label,feature_json,created_at) VALUES(?,?,?)", (title, current["feature_json"], datetime.now(timezone.utc).isoformat()))
    return jsonify(ok=bool(cursor.rowcount))


@app.delete("/api/processes/<int:process_id>")
def delete_process(process_id):
    with connect() as db:
        row = db.execute("SELECT video_id FROM processes WHERE id=?", (process_id,)).fetchone()
        if not row:
            return jsonify(error="作業が見つかりません。"), 404
        db.execute("DELETE FROM processes WHERE id=?", (process_id,))
        rows = db.execute("SELECT id FROM processes WHERE video_id=? ORDER BY position", (row["video_id"],)).fetchall()
        for position, item in enumerate(rows, 1):
            db.execute("UPDATE processes SET position=? WHERE id=?", (position, item["id"]))
    return jsonify(ok=True)


@app.post("/api/processes/<int:process_id>/ai")
def ai_description(process_id):
    with connect() as db:
        process = db.execute("SELECT title FROM processes WHERE id=?", (process_id,)).fetchone()
    if not process:
        return jsonify(error="作業が見つかりません。"), 404
    prompt = f"製造現場の動画マニュアルです。作業『{process['title']}』の説明を日本語で1文、30文字以内で書いてください。"
    try:
        data = json.dumps({"model": "phi4", "prompt": prompt, "stream": False}).encode()
        req = Request("http://127.0.0.1:11434/api/generate", data=data, headers={"Content-Type": "application/json"})
        with urlopen(req, timeout=90) as response:
            text = json.loads(response.read())["response"].strip()[:30]
    except Exception:
        return jsonify(error="Ollamaに接続できません。Ollamaを起動して phi4 を準備してください。"), 503
    with connect() as db:
        db.execute("UPDATE processes SET description_ja=? WHERE id=?", (text, process_id))
    return jsonify(description_ja=text)


if __name__ == "__main__":
    app.run(debug=True, port=5051)
